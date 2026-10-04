"""Offline reproduction, exact-budget, workbook and portability checks."""

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import export_claude_code_results as export
import run_claude_code_experiment as serial
import run_full_paired_token_experiment as pipeline
import run_token_cost_experiment as ruter
import summarize_token_usage

ROOT = Path(__file__).resolve().parents[2]
PUBLIC = ROOT / "data/claude_code/paired_cases.csv"


def observation(index, ru=2, cc=10, function="function_0001", success=1):
    return dict(zip(export.FIELDS, (
        f"case_{index:04d}", function, "run_001", "example", "gemini-2.5-flash-nothinking", "E0433",
        success, ru, 0, ru, int(ru > 0), success, cc, 0, cc, int(cc > 0), 0, 0,
    )))


class PublicResultTests(unittest.TestCase):
    def test_full_published_population_reproduces_counts(self):
        rows = export.load_cases(PUBLIC)
        results = export.summarize(rows)
        self.assertEqual(len(rows), 1858)
        self.assertEqual(len({r['function_id'] for r in rows}), 566)
        self.assertEqual(len({r['crate'] for r in rows}), 9)
        self.assertEqual([r['repaired'] for r in results], [872, 251, 600, 1297])
        self.assertEqual([r['salvaged'] for r in results], [375, 171, 320, 495])
        self.assertEqual(sum(r['ruter_total_tokens'] for r in rows), 8663677)
        self.assertEqual(sum(r['claude_total_tokens'] for r in rows), 218149284)
        self.assertAlmostEqual(results[-1]['full_run_token_ratio'], 25.179757278578137)

    def test_exact_inclusive_budget_boundary_and_failure(self):
        rows = [observation(1, cc=10), observation(2, cc=11), observation(3, cc=10, success=0)]
        export.set_budget_flags(rows)
        self.assertEqual([r['claude_success_5x'] for r in rows], [1, 0, 0])
        export.validate_rows(rows)

    def test_fractional_mean_is_not_rounded_before_comparison(self):
        rows = [observation(1, ru=1, cc=7), observation(2, ru=2, cc=8)]
        export.set_budget_flags(rows)
        self.assertEqual([r['claude_success_5x'] for r in rows], [1, 0])

    def test_function_salvage_is_any_success_not_attempt_sum(self):
        rows = [observation(1), observation(2), observation(3, function='function_0002', success=0)]
        export.set_budget_flags(rows)
        result = export.summarize(rows)[1]
        self.assertEqual((result['repaired'], result['salvaged'], result['function_units']), (2, 1, 2))

    def test_valid_rule_only_zero_remains_in_mean_and_denominator(self):
        rows = [observation(1, ru=0), observation(2, ru=4)]
        export.set_budget_flags(rows)
        result = export.summarize(rows)[0]
        self.assertEqual((result['mean_full_run_tokens'], result['attempts']), (2, 2))

    def test_duplicate_case_and_extra_private_column_are_rejected(self):
        row = observation(1)
        export.set_budget_flags([row])
        with self.assertRaises(ValueError):
            export.validate_rows([row, dict(row)])
        row['private_path'] = 'must-not-export'
        with self.assertRaises(ValueError):
            export.validate_rows([row])

    def test_inconsistent_component_totals_are_rejected(self):
        row = observation(1)
        export.set_budget_flags([row])
        row['claude_input_tokens'] = 11
        with self.assertRaises(ValueError):
            export.validate_rows([row])

    def test_csv_formula_injection_is_rejected(self):
        row = observation(1)
        export.set_budget_flags([row])
        row['crate'] = '=EXEC("unsafe")'
        with self.assertRaises(ValueError):
            export.validate_rows([row])

    def test_tampered_budget_flags_are_rejected(self):
        row = observation(1)
        export.set_budget_flags([row])
        row['claude_success_5x'] = 0
        with self.assertRaises(ValueError):
            export.validate_rows([row])

    def test_csv_export_uses_git_friendly_line_endings(self):
        row = observation(1)
        export.set_budget_flags([row])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'cases.csv'
            export.write_csv(path, export.FIELDS, [row])
            self.assertNotIn(b'\r', path.read_bytes())
            self.assertEqual(export.load_cases(path), [row])

    def test_exporter_uses_allowlist_and_separates_generation_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            ru, cc = [], []
            for i, run in enumerate(('private-generation-run-a', 'private-generation-run-b'), 1):
                common = {'attempt_uid': f'private-uid-{i}', 'run_id': run, 'node_id': 'same-function',
                          'case_id': f'private-case-{i}', 'crate': 'example', 'generation_model': 'gemini-2.5-flash-nothinking',
                          'error_codes': ['E0433']}
                ru.append({**common, 'analysis_included': True, 'sampling_weight': 1, 'strict_success': True,
                           'request_count': 1, 'input_tokens': 2, 'output_tokens': 0, 'total_tokens': 2})
                value = {**common, 'eligible': True, 'strict_success': True, 'model_valid': True,
                         'command': ['private-do-not-export'], 'account': 'private-do-not-export',
                         'usage': {'provider_usage_complete': True, 'usage_complete': True, 'request_count': 1,
                                   'provider_input_tokens': 10, 'provider_output_tokens': 0, 'provider_total_tokens': 10}}
                folder = root / f'cases/case_{i}'
                folder.mkdir(parents=True)
                (folder / 'case_result.json').write_text(json.dumps(value))
                cc.append(common)
            (root / 'experiment_manifest.json').write_text(json.dumps({'cases': cc}))
            with patch.object(summarize_token_usage, 'collect_case_rows', return_value=ru):
                rows = export.collect(root, root)
                self.assertEqual(len({r['function_id'] for r in rows}), 2)
                self.assertNotIn('private-', json.dumps(rows))
                ru[0]['analysis_included'] = False
                with self.assertRaises(ValueError):
                    export.collect(root, root)

    def test_workbook_formulas_caches_and_anonymous_metadata(self):
        from openpyxl import load_workbook
        path = ROOT / 'data/claude_code/claude_code_results.xlsx'
        calculated = load_workbook(path, data_only=True)
        try:
            calc = calculated['Calculations']
            self.assertEqual([calc[f'B{i}'].value for i in (2, 3, 4, 5, 14, 15, 16, 17, 22, 23, 24, 25)],
                             [1858, 566, 8663677, 218149284, 872, 251, 600, 1297, 375, 171, 320, 495])
            self.assertEqual(sum(calculated['Cases'].cell(i, 17).value for i in range(2, 1860)), 251)
            self.assertEqual(sum(calculated['Functions'].cell(i, 9).value for i in range(2, 568)), 320)
            self.assertEqual(calculated.properties.creator, 'Anonymous')
            self.assertEqual(calculated.properties.lastModifiedBy, 'Anonymous')
            self.assertEqual(calculated.properties.modified.year, 2000)
        finally:
            calculated.close()
        formulas = load_workbook(path, data_only=False)
        try:
            self.assertEqual(formulas['Calculations']['B6'].value, '=B4/B2')
            self.assertIn('Calculations!$B$10', formulas['Cases']['Q2'].value)
            self.assertIn('COUNTIFS', formulas['Functions']['H2'].value)
        finally:
            formulas.close()
        with ZipFile(path) as archive:
            contents = b'\n'.join(archive.read(n) for n in archive.namelist() if n.endswith('.xml'))
            for forbidden in (b'/home/', b'/Users/', b'Authorization:', b'private-do-not-export'):
                self.assertNotIn(forbidden, contents)
            self.assertFalse(any('externalLink' in n for n in archive.namelist()))


class PortableConfigurationTests(unittest.TestCase):
    def test_default_paths_are_inside_standalone_repository(self):
        for path in (ruter.DEFAULT_MANIFEST, ruter.DEFAULT_CRATES_ROOT, serial.DEFAULT_PAIRED_MANIFEST,
                     serial.DEFAULT_SETTINGS, serial.DEFAULT_PROMPT, serial.DEFAULT_VERIFY_HOOK,
                     pipeline.DEFAULT_MANIFEST, pipeline.DEFAULT_CRATES_ROOT, pipeline.DEFAULT_CONTROL_ROOT):
            self.assertTrue(path.is_relative_to(ROOT), path)

    def test_fresh_pipeline_does_not_reuse_private_historical_outputs(self):
        args = pipeline.build_parser().parse_args([])
        self.assertIsNone(args.reuse_ruter_results_from)
        self.assertIsNone(args.reuse_claude_results_from)

    def test_packaged_prompt_hook_and_settings_exist_without_credentials(self):
        settings = json.loads(serial.DEFAULT_SETTINGS.read_text())
        self.assertEqual(settings['env']['ANTHROPIC_AUTH_TOKEN'], 'local-token-meter-proxy')
        self.assertIn('{failed_test}', serial.DEFAULT_PROMPT.read_text())
        self.assertTrue(serial.DEFAULT_VERIFY_HOOK.is_file())

    def test_pipeline_forwards_input_root_and_isolates_targets(self):
        args = pipeline.build_parser().parse_args(['--defer-infrastructure-retries', '--claude-workers', '2'])
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch.object(pipeline, 'run_deferred_pipeline', return_value=True) as mocked:
                self.assertTrue(pipeline.run_pipeline(args, 1858, 'https://example.invalid',
                                root / 'ruter', root / 'claude', root / 'control', {}, {}, {}))
            ru, cc = mocked.call_args.args[2:4]
            self.assertEqual(ru[ru.index('--clean-crates-root') + 1], str(args.clean_crates_root.resolve()))
            self.assertEqual(ru[ru.index('--cargo-target-root') + 1], str(root / 'ruter/_cargo_target'))
            self.assertEqual(cc[cc.index('--cargo-target-root') + 1], str(root / 'claude/_cargo_target'))
            self.assertNotIn('--reuse-results-from', ru)
            self.assertNotIn('--reuse-results-from', cc)
            self.assertEqual(cc[cc.index('--workers') + 1], '2')


if __name__ == '__main__':
    unittest.main()

#!/usr/bin/env python3
"""Export paired results via a strict allowlist, or recompute the public tables.

No API calls. Raw prompts, paths, identities, timestamps and provider IDs are
never exported. Workbook formulas have cached values for non-Excel readers.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

FIELDS = (
    "case_id", "function_id", "run_id", "crate", "generation_model", "error_codes",
    "ruter_success", "ruter_input_tokens", "ruter_output_tokens", "ruter_total_tokens",
    "ruter_requests", "claude_success", "claude_input_tokens", "claude_output_tokens",
    "claude_total_tokens", "claude_requests", "claude_success_5x", "claude_success_10x",
)
SUMMARY_FIELDS = (
    "method", "budget_multiple", "attempts", "repaired", "repair_rate",
    "function_units", "salvaged", "salvage_rate", "mean_full_run_tokens",
    "full_run_token_ratio", "token_cap",
)
STRING_FIELDS = set(FIELDS[:6])
XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def flag(value):
    if value is True or value == "True":
        return True
    if value is False or value == "False":
        return False
    raise ValueError("missing or invalid boolean in raw result")


def nonnegative(value):
    if isinstance(value, bool):
        raise ValueError("a boolean is not a token count")
    number = int(value)
    if number < 0 or str(number) != str(value):
        raise ValueError("invalid nonnegative integer")
    return number


def collect(ruter_root: Path, claude_root: Path):
    from summarize_token_usage import collect_case_rows

    ruter = collect_case_rows(ruter_root)
    manifest = read_json(claude_root / "experiment_manifest.json")
    cases = manifest["cases"]
    uids = [str(row["attempt_uid"]) for row in cases]
    if not cases or len(set(uids)) != len(uids):
        raise ValueError("empty or duplicate paired population")
    ru_map = {str(row["attempt_uid"]): row for row in ruter}
    if len(ru_map) != len(ruter) or set(ru_map) != set(uids):
        raise ValueError("RuTeR and Claude populations differ")
    cc_map = {}
    for path in claude_root.glob("cases/*/case_result.json"):
        row = read_json(path)
        uid = str(row["attempt_uid"])
        if uid in cc_map:
            raise ValueError("duplicate Claude result")
        cc_map[uid] = row
    if set(cc_map) != set(uids):
        raise ValueError("incomplete or mismatched Claude population")
    if any(float(row["sampling_weight"]) != 1 for row in ruter):
        raise ValueError("this exporter requires a full, unweighted population")
    run_ids = {run: f"run_{i:03d}" for i, run in enumerate(sorted({str(c['run_id']) for c in cases}), 1)}
    functions = sorted({(str(c["run_id"]), str(c["node_id"])) for c in cases})
    function_ids = {key: f"function_{i:04d}" for i, key in enumerate(functions, 1)}
    rows = []
    for index, case in enumerate(sorted(cases, key=lambda c: c["case_id"]), 1):
        ru, cc = ru_map[str(case["attempt_uid"])], cc_map[str(case["attempt_uid"])]
        usage = cc.get("usage") or {}
        if not (flag(ru["analysis_included"]) and flag(cc.get("eligible"))
                and flag(cc.get("model_valid")) and flag(usage.get("provider_usage_complete"))
                and flag(usage.get("usage_complete"))):
            raise ValueError("ineligible case, wrong model or incomplete provider usage")
        if (str(cc["run_id"]), str(cc["node_id"]), str(cc["crate"])) != (
                str(case["run_id"]), str(case["node_id"]), str(case["crate"])):
            raise ValueError("Claude metadata does not match the paired manifest")
        row = {
            "case_id": f"case_{index:04d}",
            "function_id": function_ids[(str(case["run_id"]), str(case["node_id"]))],
            "run_id": run_ids[str(case["run_id"])],
            "crate": str(case["crate"]), "generation_model": str(case["generation_model"]),
            "error_codes": ",".join(sorted(set(case.get("error_codes") or []))),
            "ruter_success": int(flag(ru["strict_success"])),
            "claude_success": int(flag(cc["strict_success"])),
            "ruter_requests": nonnegative(ru["request_count"]),
            "claude_requests": nonnegative(usage["request_count"]),
        }
        for component in ("input_tokens", "output_tokens", "total_tokens"):
            row[f"ruter_{component}"] = nonnegative(ru[component])
            row[f"claude_{component}"] = nonnegative(usage[f"provider_{component}"])
        rows.append(row)
    set_budget_flags(rows)
    validate_rows(rows)
    return rows


def set_budget_flags(rows):
    total = sum(row["ruter_total_tokens"] for row in rows)
    for row in rows:
        for multiple in (5, 10):
            row[f"claude_success_{multiple}x"] = int(
                row["claude_success"] == 1 and row["claude_total_tokens"] * len(rows) <= multiple * total
            )


def validate_rows(rows):
    if not rows or len({r["case_id"] for r in rows}) != len(rows):
        raise ValueError("empty table or duplicate case ID")
    function_metadata = {}
    for row in rows:
        if set(row) != set(FIELDS):
            raise ValueError("unexpected or missing release columns")
        for field, pattern in (("case_id", r"case_\d+"), ("function_id", r"function_\d+"), ("run_id", r"run_\d+"),
                               ("crate", r"[A-Za-z0-9_-]+"), ("generation_model", r"[A-Za-z0-9_.-]+"),
                               ("error_codes", r"(?:E\d{4}(?:,E\d{4})*)?")):
            if not re.fullmatch(pattern, row[field]):
                raise ValueError(f"invalid public {field}")
        metadata = (row["run_id"], row["crate"], row["generation_model"])
        old = function_metadata.setdefault(row["function_id"], metadata)
        if old != metadata:
            raise ValueError("function crosses generation runs or crates")
        for field in set(FIELDS) - STRING_FIELDS:
            nonnegative(row[field])
            if "success" in field and row[field] not in (0, 1):
                raise ValueError("invalid success flag")
        for method in ("ruter", "claude"):
            if row[f"{method}_input_tokens"] + row[f"{method}_output_tokens"] != row[f"{method}_total_tokens"]:
                raise ValueError("provider total does not equal input plus output")
            if row[f"{method}_requests"] == 0 and row[f"{method}_total_tokens"] != 0:
                raise ValueError("nonzero tokens without a request")
    expected = [dict(r) for r in rows]
    set_budget_flags(expected)
    if any(a != b for a, b in zip(rows, expected)):
        raise ValueError("budget flags do not match exact, unrounded mean")


def load_cases(path):
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != list(FIELDS):
            raise ValueError("unexpected case CSV schema")
        rows = [{key: value if key in STRING_FIELDS else nonnegative(value)
                 for key, value in row.items()} for row in reader]
    validate_rows(rows)
    return rows


def summarize(rows):
    validate_rows(rows)
    n = len(rows)
    functions = {r["function_id"] for r in rows}
    ru_tokens = sum(r["ruter_total_tokens"] for r in rows)
    cc_tokens = sum(r["claude_total_tokens"] for r in rows)
    if not ru_tokens:
        raise ValueError("RuTeR mean is zero; token ratios are undefined")
    result = []
    for method, multiple, field in (
            ("RuTeR", "", "ruter_success"), ("Claude Code", 5, "claude_success_5x"),
            ("Claude Code", 10, "claude_success_10x"), ("Claude Code", "", "claude_success")):
        repaired = sum(r[field] for r in rows)
        salvaged = len({r["function_id"] for r in rows if r[field]})
        tokens = ru_tokens if method == "RuTeR" else cc_tokens
        result.append(dict(zip(SUMMARY_FIELDS, (
            method, multiple, n, repaired, repaired / n, len(functions), salvaged,
            salvaged / len(functions), tokens / n, tokens / ru_tokens,
            multiple * ru_tokens / n if multiple else "",
        ))))
    return result


def write_csv(path, fields, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def cache_formulas(path, cached):
    # openpyxl does not write formula caches. Cache only our known formula cells;
    # Excel is still instructed to recalculate if a reviewer changes the inputs.
    payload = io.BytesIO()
    with ZipFile(path) as source, ZipFile(payload, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl overwrites modified on save; remove that fingerprint.
                data = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>", data,
                )
            if info.filename in cached:
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{XML_NS}}}c"):
                    if cell.get("r") not in cached[info.filename]:
                        continue
                    cell.attrib.pop("t", None)
                    value = cell.find(f"{{{XML_NS}}}v")
                    if value is None:
                        value = ET.SubElement(cell, f"{{{XML_NS}}}v")
                    value.text = str(cached[info.filename][cell.get("r")])
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            # Neutral ZIP timestamps; no source file metadata or local paths.
            info.date_time = (2000, 1, 1, 0, 0, 0)
            target.writestr(info, data)
    path.write_bytes(payload.getvalue())


def write_workbook(path, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.workbook.properties import CalcProperties
    except ImportError as error:
        raise SystemExit("Excel export requires openpyxl; see experiments/requirements-token-cost.txt") from error
    summary = summarize(rows)
    wb = Workbook()
    wb.properties.creator = "Anonymous"
    wb.properties.lastModifiedBy = "Anonymous"
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    readme = wb.active
    readme.title = "README"
    for entry in (
        ("Item", "Definition"),
        ("Scope", "Supplementary full paired rerun; original submitted-paper aggregates remain unchanged."),
        ("Population", f"{len(rows)} failed generation attempts; {len({r['function_id'] for r in rows})} generation-run/function units."),
        ("Model", "gemini-2.5-flash-nothinking; provider alias gemini-2.5-flash; same gateway and API account for both methods."),
        ("Gateway", "https://4sapi.org (OpenAI chat/completions for RuTeR; Anthropic messages for Claude Code)."),
        ("Claude Code", "2.1.281; max 20 turns; case timeout 900s; Cargo timeout 600s; two isolated workers in the full run."),
        ("RuTeR", "topk=3; max three LLM rounds per unresolved function; case timeout 900s; LLM timeout 120s."),
        ("Tokens", "Provider input + output for all requests in the final valid trial, including ordinary failures and valid rule-only zero-token cases."),
        ("Cache", "Claude primary cost is billing_usage.openai_usage.total_tokens; do not add cache-read counts again."),
        ("Strict success", "RuTeR: passing patch verification, zero errors/unresolved test functions. Claude: passing final compile, zero errors, scope/integrity guards; not semantic equivalence."),
        ("Function grouping", "function_id is a pseudonym for (generation run_id, node_id). Different runs remain distinct."),
        ("Salvage", "A function is salvaged iff at least one of its failed generation attempts has an eligible strict repair."),
        ("Budget", "Post-hoc full-run cutoff, NOT an online early-stop experiment. Per attempt: Claude total <= multiple * full-population RuTeR mean."),
        ("Denominators", "All paired attempts and all function units stay in the denominator; over-budget repairs count as failures."),
        ("Retry scope", "Final provider-complete trials only; preflight and archived infrastructure retries excluded. Ordinary method failures are never rerun for success."),
        ("Calculations", "Editable multipliers are in Calculations B10/B11. Cases Q/R, Functions and Calculations recompute the results; cached values are supplied."),
        ("Summary", "Verified snapshot. mean_full_run_tokens and full_run_token_ratio are untruncated costs, also on the post-hoc budget rows."),
        ("Privacy", "Only allowlisted counts, flags, public crate/model names and pseudonymous IDs. No prompts, account IDs, machine paths or credentials."),
        ("Rerunning", "Raw failed tests and clean crate snapshots are separate inputs, not reconstructed by this compact workbook. See experiments/CLAUDE_CODE_EXPERIMENT.md."),
    ):
        readme.append(entry)
    sheet = wb.create_sheet("Summary")
    sheet.append(SUMMARY_FIELDS)
    for row in summary:
        sheet.append([row[f] for f in SUMMARY_FIELDS])
    calc = wb.create_sheet("Calculations")
    calc.append(("Metric", "Value"))
    calc_rows = (
        ("Attempts", len(rows)), ("Functions", len({r['function_id'] for r in rows})),
        ("RuTeR total tokens", sum(r['ruter_total_tokens'] for r in rows)),
        ("Claude total tokens", sum(r['claude_total_tokens'] for r in rows)),
        ("RuTeR mean tokens", summary[0]['mean_full_run_tokens']),
        ("Claude mean tokens", summary[3]['mean_full_run_tokens']),
        ("Claude/RuTeR ratio", summary[3]['full_run_token_ratio']),
        ("", ""), ("First budget multiplier", 5), ("Second budget multiplier", 10),
        ("First exact token cap", summary[1]['token_cap']),
        ("Second exact token cap", summary[2]['token_cap']),
        ("RuTeR repairs", summary[0]['repaired']), ("First-budget Claude repairs", summary[1]['repaired']),
        ("Second-budget Claude repairs", summary[2]['repaired']), ("Unlimited Claude repairs", summary[3]['repaired']),
        ("RuTeR repair rate", summary[0]['repair_rate']), ("First-budget Claude repair rate", summary[1]['repair_rate']),
        ("Second-budget Claude repair rate", summary[2]['repair_rate']), ("Unlimited Claude repair rate", summary[3]['repair_rate']),
        ("RuTeR salvaged", summary[0]['salvaged']), ("First-budget Claude salvaged", summary[1]['salvaged']),
        ("Second-budget Claude salvaged", summary[2]['salvaged']), ("Unlimited Claude salvaged", summary[3]['salvaged']),
        ("RuTeR salvage rate", summary[0]['salvage_rate']), ("First-budget Claude salvage rate", summary[1]['salvage_rate']),
        ("Second-budget Claude salvage rate", summary[2]['salvage_rate']), ("Unlimited Claude salvage rate", summary[3]['salvage_rate']),
    )
    for row in calc_rows:
        calc.append(row)
    cases = wb.create_sheet("Cases")
    cases.append(FIELDS)
    for row in rows:
        cases.append([row[f] for f in FIELDS])
    functions = wb.create_sheet("Functions")
    functions.append(("function_id", "run_id", "crate", "generation_model", "attempts", "ruter_salvaged", "claude_salvaged", "claude_salvaged_5x", "claude_salvaged_10x"))
    grouped = defaultdict(list)
    for row in rows:
        grouped[row['function_id']].append(row)
    function_values = []
    for uid, group in sorted(grouped.items()):
        first = group[0]
        values = [uid, first['run_id'], first['crate'], first['generation_model'], len(group)]
        values.extend(int(any(r[f] for r in group)) for f in ('ruter_success', 'claude_success', 'claude_success_5x', 'claude_success_10x'))
        functions.append(values)
        function_values.append(values)
    last, flast = len(rows) + 1, len(grouped) + 1
    cached = {f"xl/worksheets/sheet{i}.xml": {} for i in (3, 4, 5)}
    calc_forms = {
        2: f"COUNTA(Cases!A2:A{last})", 3: f"COUNTA(Functions!A2:A{flast})",
        4: f"SUM(Cases!J2:J{last})", 5: f"SUM(Cases!O2:O{last})",
        6: "B4/B2", 7: "B5/B2", 8: "B5/B4", 12: "B10*B6", 13: "B11*B6",
        14: f"SUM(Cases!G2:G{last})", 15: f"SUM(Cases!Q2:Q{last})",
        16: f"SUM(Cases!R2:R{last})", 17: f"SUM(Cases!L2:L{last})",
        18: "B14/B2", 19: "B15/B2", 20: "B16/B2", 21: "B17/B2",
        22: f"SUM(Functions!F2:F{flast})", 23: f"SUM(Functions!H2:H{flast})",
        24: f"SUM(Functions!I2:I{flast})", 25: f"SUM(Functions!G2:G{flast})",
        26: "B22/B3", 27: "B23/B3", 28: "B24/B3", 29: "B25/B3",
    }
    for row, formula in calc_forms.items():
        address = f"B{row}"
        cached['xl/worksheets/sheet3.xml'][address] = calc[address].value
        calc[address] = "=" + formula
    for i, row in enumerate(rows, 2):
        for column, multiple, parameter in (('Q', 5, 10), ('R', 10, 11)):
            address = f"{column}{i}"
            cached['xl/worksheets/sheet4.xml'][address] = row[f'claude_success_{multiple}x']
            cases[address] = f'=IF(AND(L{i}=1,O{i}*Calculations!$B$2<=Calculations!$B${parameter}*Calculations!$B$4),1,0)'
    for i, values in enumerate(function_values, 2):
        cached['xl/worksheets/sheet5.xml'][f'E{i}'] = values[4]
        functions[f'E{i}'] = f'=COUNTIF(Cases!$B$2:$B${last},A{i})'
        for col, source, value in zip(('F', 'G', 'H', 'I'), ('G', 'L', 'Q', 'R'), values[5:]):
            cached['xl/worksheets/sheet5.xml'][f'{col}{i}'] = value
            functions[f'{col}{i}'] = f'=IF(COUNTIFS(Cases!$B$2:$B${last},A{i},Cases!${source}$2:${source}${last},1)>0,1,0)'
    for ws in wb:
        ws.freeze_panes = 'A2'
        if ws.title != 'README':
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='264653')
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(45, max(14, len(str(column[0].value)) + 3))
    readme.column_dimensions['B'].width = 120
    for row in sheet.iter_rows(min_row=2):
        for index in (4, 7):
            row[index].number_format = '0.00%'
        for index in (8, 9, 10):
            row[index].number_format = '0.00'
    for row in (18, 19, 20, 21, 26, 27, 28, 29):
        calc[f'B{row}'].number_format = '0.00%'
    wb.save(path)
    cache_formulas(path, cached)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ruter-root', type=Path)
    parser.add_argument('--claude-root', type=Path)
    parser.add_argument('--cases-csv', type=Path, help='Recompute from published compact data; no raw logs required')
    parser.add_argument('--out', type=Path, help='Output directory; omit to print results only')
    args = parser.parse_args()
    if args.cases_csv:
        if args.ruter_root or args.claude_root:
            parser.error('choose either public CSV or both raw experiment roots')
        rows = load_cases(args.cases_csv)
    elif args.ruter_root and args.claude_root:
        rows = collect(args.ruter_root, args.claude_root)
    else:
        parser.error('supply --cases-csv, or both --ruter-root and --claude-root')
    summary = summarize(rows)
    print(json.dumps(summary, indent=2))
    if args.out:
        args.out.mkdir(parents=True, exist_ok=True)
        write_workbook(args.out / 'claude_code_results.xlsx', rows)
        write_csv(args.out / 'paired_cases.csv', FIELDS, rows)
        write_csv(args.out / 'summary.csv', SUMMARY_FIELDS, summary)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
